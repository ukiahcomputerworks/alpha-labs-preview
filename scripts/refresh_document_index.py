"""Refresh the public Alpha Labs document search index without changing site pages.

Only files linked from the retained public routes are eligible. PDF pages with
embedded text are extracted directly; image-only PDF pages and service-sheet
images are OCR'd. Previous successful text is retained and marked stale if a
temporary fetch or extraction failure occurs.
"""

from __future__ import annotations

import argparse
import hashlib
import html
from html.parser import HTMLParser
import io
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import time
from urllib.error import HTTPError
from urllib.parse import quote, unquote, urljoin, urlsplit, urlunsplit
from urllib.request import HTTPRedirectHandler, Request, build_opener


ROOT = Path(__file__).resolve().parent.parent
ALLOWED_HOSTS = {"alpha-labs.com", "www.alpha-labs.com"}
DOCUMENT_SUFFIXES = {".pdf", ".xlsx"}
IMAGE_SUFFIXES = {".webp", ".png", ".jpg", ".jpeg"}
MAX_BYTES = 30 * 1024 * 1024
MAX_PDF_PAGES = 200
USER_AGENT = "Mozilla/5.0 (compatible; AlphaLabsPreviewDocumentIndexer/1.0)"
TOKEN_PATTERN = re.compile(r"[^\W_]+", re.UNICODE)


class FirstPartyRedirects(HTTPRedirectHandler):
    def redirect_request(self, request, fp, code, msg, headers, newurl):
        if urlsplit(newurl).hostname not in ALLOWED_HOSTS:
            raise ValueError(f"Refusing off-site redirect: {newurl}")
        return super().redirect_request(request, fp, code, msg, headers, newurl)


OPENER = build_opener(FirstPartyRedirects())


def compact(value: str) -> str:
    return re.sub(r"\s+", " ", html.unescape(value)).strip()


def tokenize(value: str) -> list[str]:
    import unicodedata

    value = unicodedata.normalize("NFKD", value).lower()
    value = "".join(char for char in value if not unicodedata.combining(char))
    return TOKEN_PATTERN.findall(value)


def normalize_url(raw: str, base: str, *, image: bool = False) -> str | None:
    candidate = urljoin(base, html.unescape(raw).strip())
    parts = urlsplit(candidate)
    if parts.hostname not in ALLOWED_HOSTS:
        return None
    suffix = Path(unquote(parts.path)).suffix.lower()
    if suffix not in (IMAGE_SUFFIXES if image else DOCUMENT_SUFFIXES):
        return None
    safe_path = quote(unquote(parts.path), safe="/._~-()")
    return urlunsplit(("https", "www.alpha-labs.com", safe_path, parts.query, ""))


def filename_title(url: str) -> str:
    stem = Path(unquote(urlsplit(url).path)).stem
    return compact(re.sub(r"[_-]+", " ", stem)) or "Document"


class PublicContentLinks(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.main_depth = 0
        self.current_anchor: dict | None = None
        self.links: list[tuple[str, str, str]] = []

    def handle_starttag(self, tag, attrs):
        attributes = dict(attrs)
        if tag == "main":
            self.main_depth += 1
        if not self.main_depth:
            return
        if tag == "a" and attributes.get("href"):
            self.current_anchor = {"href": attributes["href"], "text": []}
        if tag == "area" and attributes.get("href"):
            self.links.append((attributes["href"], compact(attributes.get("alt") or attributes.get("title") or ""), "document"))
        if tag == "img" and "program-work-img" in attributes.get("class", "").split():
            source = attributes.get("src")
            if source:
                self.links.append((source, compact(attributes.get("alt", "")), "image"))

    def handle_data(self, data):
        if self.current_anchor is not None:
            self.current_anchor["text"].append(data)

    def handle_endtag(self, tag):
        if tag == "a" and self.current_anchor is not None:
            self.links.append((self.current_anchor["href"], compact(" ".join(self.current_anchor["text"])), "document"))
            self.current_anchor = None
        if tag == "main" and self.main_depth:
            self.main_depth -= 1


def page_href(route: str) -> str:
    return "" if route == "/" else route.lstrip("/") + ("" if route.endswith(".html") else "/")


def add_links(inventory: dict, markup: str, source_url: str, route: str):
    parser = PublicContentLinks()
    parser.feed(markup)
    for raw, label, kind in parser.links:
        url = normalize_url(raw, source_url, image=kind == "image")
        if not url:
            continue
        item = inventory.setdefault(url, {"href": url, "title": "", "kind": kind, "sourcePages": set()})
        item["sourcePages"].add(route)
        title = label if len(label) >= 8 and label.lower() not in {"click here", "download pdf", "download"} else filename_title(url)
        if kind == "image":
            item["href"] = page_href(route)
            title = f"{title} (service sheet)"
        if len(title) > len(item["title"]):
            item["title"] = title[:160]


def request_bytes(url: str, *, etag: str = "", last_modified: str = "", max_bytes: int = MAX_BYTES):
    headers = {"User-Agent": USER_AGENT, "Accept": "*/*"}
    if etag:
        headers["If-None-Match"] = etag
    if last_modified:
        headers["If-Modified-Since"] = last_modified
    request = Request(url, headers=headers)
    try:
        response = OPENER.open(request, timeout=25)
    except HTTPError as error:
        if error.code == 304:
            return None, etag, last_modified
        raise
    with response:
        if urlsplit(response.geturl()).hostname not in ALLOWED_HOSTS:
            raise ValueError("Response left the allowed Alpha Labs host")
        length = int(response.headers.get("Content-Length") or 0)
        if length > max_bytes:
            raise ValueError(f"File exceeds the {max_bytes // 1024 // 1024} MiB limit")
        chunks = []
        count = 0
        while True:
            part = response.read(1024 * 1024)
            if not part:
                break
            count += len(part)
            if count > max_bytes:
                raise ValueError(f"File exceeds the {max_bytes // 1024 // 1024} MiB limit")
            chunks.append(part)
        return b"".join(chunks), response.headers.get("ETag", ""), response.headers.get("Last-Modified", "")


def ocr_png(png: Path) -> str:
    if not shutil.which("tesseract"):
        raise RuntimeError("Tesseract OCR is unavailable")
    result = subprocess.run(["tesseract", str(png), "stdout", "-l", "eng"], capture_output=True, text=True, timeout=120)
    if result.returncode:
        raise RuntimeError(f"Tesseract failed: {result.stderr[:180]}")
    return compact(result.stdout)


def extract_pdf(data: bytes, temp: Path) -> tuple[str, str, int]:
    from pypdf import PdfReader

    pdf = temp / "document.pdf"
    pdf.write_bytes(data)
    reader = PdfReader(io.BytesIO(data), strict=False)
    if reader.is_encrypted:
        if reader.decrypt("") == 0:
            raise ValueError("Encrypted PDF cannot be read")
    count = len(reader.pages)
    if count > MAX_PDF_PAGES:
        raise ValueError(f"PDF has {count} pages; limit is {MAX_PDF_PAGES}")
    sections = []
    ocr_pages = 0
    for number, page in enumerate(reader.pages, start=1):
        native = compact(page.extract_text() or "")
        if len(re.sub(r"\W", "", native)) < 80:
            if not shutil.which("pdftoppm"):
                raise RuntimeError("Poppler rendering is unavailable for OCR")
            prefix = temp / f"page-{number}"
            command = ["pdftoppm", "-f", str(number), "-l", str(number), "-r", "175", "-singlefile", "-png", str(pdf), str(prefix)]
            result = subprocess.run(command, capture_output=True, text=True, timeout=120)
            if result.returncode:
                raise RuntimeError(f"PDF page render failed: {result.stderr[:180]}")
            native = ocr_png(prefix.with_suffix(".png"))
            ocr_pages += 1
        sections.append(native)
    text = compact(" ".join(sections))
    if not text:
        raise ValueError("No searchable text was extracted")
    mode = "ocr" if ocr_pages == count else "mixed" if ocr_pages else "embedded-text"
    return text, mode, count


def extract_image(data: bytes, temp: Path) -> tuple[str, str, int]:
    from PIL import Image

    image = Image.open(io.BytesIO(data))
    png = temp / "sheet.png"
    image.convert("RGB").save(png)
    text = ocr_png(png)
    if not text:
        raise ValueError("Service sheet OCR produced no text")
    return text, "ocr", 1


def extract_spreadsheet(data: bytes) -> tuple[str, str, int]:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sections = []
    for sheet in book:
        sections.append(sheet.title)
        for row in sheet.iter_rows(values_only=True):
            sections.extend(str(value) for value in row if value is not None)
    text = compact(" ".join(sections))
    if not text:
        raise ValueError("Spreadsheet has no readable cell values")
    return text, "spreadsheet-cells", len(book.sheetnames)


def build_terms(documents: list[dict]) -> dict:
    postings: dict[str, list[list[int]]] = {}
    for index, document in enumerate(documents):
        frequencies: dict[str, int] = {}
        for word in tokenize(document["title"] + " " + document["text"]):
            frequencies[word] = frequencies.get(word, 0) + 1
        for word, count in frequencies.items():
            postings.setdefault(word, []).append([index, count])
    return dict(sorted(postings.items()))


def main() -> int:
    args = argparse.ArgumentParser(description=__doc__)
    args.add_argument("--local-only", action="store_true", help="Do not discover newly linked files from the live public source pages")
    args.add_argument("--max-documents", type=int, default=0, help="Bounded diagnostic run; use with --output outside the site root")
    args.add_argument("--inventory-only", action="store_true", help="List eligible link counts without fetching files or writing an index")
    args.add_argument("--output", type=Path, default=ROOT / "document-search-index.json")
    options = args.parse_args()
    manifest = json.loads((ROOT / "mirror-manifest.json").read_text(encoding="utf-8"))
    old = json.loads(options.output.read_text(encoding="utf-8")) if options.output.exists() else {"documents": []}
    previous = {item["sourceUrl"]: item for item in old.get("documents", []) if item.get("sourceUrl")}
    inventory: dict[str, dict] = {}
    source_failures = []
    for item in manifest:
        markup = (ROOT / item["Output"].replace("\\", os.sep)).read_text(encoding="utf-8")
        add_links(inventory, markup, item["Source"], item["Route"])
        if not options.local_only:
            try:
                source, _, _ = request_bytes(item["Source"], max_bytes=5 * 1024 * 1024)
                add_links(inventory, source.decode("utf-8", errors="replace"), item["Source"], item["Route"])
            except Exception as error:
                source_failures.append({"route": item["Route"], "reason": str(error)[:180]})
    # A temporary source-page outage must not remove a document indexed earlier.
    failed_routes = {item["route"] for item in source_failures}
    for url, item in previous.items():
        if url not in inventory and failed_routes.intersection(item.get("sourcePages", [])):
            inventory[url] = {"href": item["href"], "title": item["title"], "kind": item["kind"], "sourcePages": set(item["sourcePages"])}
    if options.inventory_only:
        image_count = sum(item["kind"] == "image" for item in inventory.values())
        print(f"Eligible first-party files: {len(inventory)} ({len(inventory) - image_count} linked documents, {image_count} service-sheet images); source-page failures: {len(source_failures)}")
        return 0
    candidates = sorted(inventory.items())
    if options.max_documents:
        candidates = candidates[: options.max_documents]
    documents = []
    failures = []
    changed = 0
    for position, (url, item) in enumerate(candidates, start=1):
        before = previous.get(url)
        try:
            data, etag, modified = request_bytes(url, etag=(before or {}).get("etag", ""), last_modified=(before or {}).get("lastModified", ""))
            if data is None:
                if not before:
                    raise ValueError("Server returned Not Modified without cached text")
                text, method, pages = before["text"], before["method"], before["units"]
                digest = before["sha256"]
            else:
                digest = hashlib.sha256(data).hexdigest()
                if before and digest == before.get("sha256"):
                    text, method, pages = before["text"], before["method"], before["units"]
                else:
                    with tempfile.TemporaryDirectory(prefix="alpha-ocr-") as folder:
                        temp = Path(folder)
                        if item["kind"] == "image":
                            text, method, pages = extract_image(data, temp)
                        elif Path(urlsplit(url).path).suffix.lower() == ".xlsx":
                            text, method, pages = extract_spreadsheet(data)
                        elif data.startswith(b"%PDF"):
                            text, method, pages = extract_pdf(data, temp)
                        else:
                            raise ValueError("The response was not the expected PDF or spreadsheet")
                    changed += 1
            documents.append({
                "title": item["title"], "href": item["href"], "sourceUrl": url,
                "sourcePages": sorted(item["sourcePages"]), "kind": item["kind"],
                "text": text, "method": method, "units": pages, "sha256": digest,
                "etag": etag, "lastModified": modified,
            })
        except Exception as error:
            failures.append({"url": url, "reason": str(error)[:180]})
            if before:
                documents.append({**before, "title": item["title"], "href": item["href"], "sourcePages": sorted(item["sourcePages"]), "stale": True})
        if position % 20 == 0 or position == len(candidates):
            print(f"Checked {position}/{len(candidates)} public files; newly extracted or changed: {changed}; failures: {len(failures)}", flush=True)
        time.sleep(0.03)
    documents.sort(key=lambda entry: entry["sourceUrl"])
    result = {
        "version": 1, "documentCount": len(documents), "targetCount": len(candidates),
        "failedCount": len(failures), "sourcePageFailures": source_failures,
        "failures": failures, "documents": documents, "terms": build_terms(documents),
    }
    options.output.parent.mkdir(parents=True, exist_ok=True)
    options.output.write_text(json.dumps(result, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    print(f"Indexed {len(documents)} of {len(candidates)} public files; {changed} content extractions; {len(failures)} file failures; {len(source_failures)} source-page failures.")
    return 0 if documents or options.max_documents else 1


if __name__ == "__main__":
    sys.exit(main())
